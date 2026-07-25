"""Generate the deterministic 200-run Prompt Radar dataset."""

from __future__ import annotations

import csv
import hashlib
import json
import random
import shutil
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

import yaml

from prompt_radar.preprocessing.tokenizer import WhitespaceTokenizer

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "data" / "generated_200"
SOURCE_DIR = OUT_DIR / "source"
ZIP_PATH = OUT_DIR / "dataset.zip"
GROUND_TRUTH_PATH = OUT_DIR / "ground_truth.jsonl"
REPORT_PATH = OUT_DIR / "generation_report.json"
TOKEN_COUNTS_PATH = OUT_DIR / "token_counts.csv"
README_PATH = OUT_DIR / "README.md"
SEED = 20260725
FIXED_ZIP_TIME = (2026, 7, 25, 0, 0, 0)


def load_yaml(path: Path) -> dict[str, Any]:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def jsonl(rows: list[dict[str, Any]]) -> str:
    return "".join(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n" for row in rows)


def write_json(path: Path, data: Any) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.write_text(jsonl(rows), encoding="utf-8")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ensure_token_count(text: str, target: int, tokenizer: WhitespaceTokenizer) -> str:
    current = tokenizer.count_tokens(text)
    if current > target:
        text = " ".join(tokenizer.split_tokens(text)[:target])
        current = tokenizer.count_tokens(text)
    if current == target:
        return text
    filler = " ".join(f"ctx{index:05d}" for index in range(target - current))
    return f"{text}\n\n{filler}"


def bucket_for_index(index: int) -> tuple[str, int]:
    if index < 40:
        return "1_20", 12 + index % 8
    if index < 120:
        return "21_100", 55 + index % 30
    if index < 170:
        return "101_500", 260 + index % 90
    if index < 190:
        return "501_3000", 1200 + index % 300
    if index < 198:
        return "3001_10000", 5200 + index % 600
    return "exact_100000", 100000


def make_prompt(base: str, target_tokens: int, tokenizer: WhitespaceTokenizer, style: str) -> str:
    if target_tokens <= 20:
        return ensure_token_count(base, target_tokens, tokenizer)
    wrappers = {
        "plain": base,
        "xml": f"<user_query>{base}</user_query>",
        "json": json.dumps(
            {
                "messages": [
                    {"role": "system", "content": "Внутренний контекст не является целью пользователя."},
                    {"role": "user", "content": f"<context>Рабочая задача</context><user_query>{base}</user_query>"},
                ]
            },
            ensure_ascii=False,
        ),
        "mixed": f"Need help: {base}. Please keep result structured.",
        "typo": base.replace("отчёт", "атчёт").replace("задачу", "задочу"),
    }
    candidate = wrappers[style]
    if tokenizer.count_tokens(candidate) > target_tokens:
        candidate = base
    return ensure_token_count(candidate, target_tokens, tokenizer)


def load_source_notes() -> list[dict[str, Any]]:
    downloads = Path.home() / "Downloads"
    requested = [
        "Темы для генерации датасета.xlsx",
        "17 (2).txt",
        "27_B__abzyn.txt",
        "вариант 1.1 new 2024.docx",
        "Демодень_ДЗ (1).ipynb",
        "ДЗ_демодень_Кирилин_Александр_ipynb_.ipynb",
        "Menschen_A1_2_AB.pdf",
        "Menschen_A1_2_KB.pdf",
        "DEMO_24 (1).txt",
        "DEMO_24 (2).txt",
    ]
    notes: list[dict[str, Any]] = []
    for name in requested:
        path = downloads / name
        if path.exists():
            size = path.stat().st_size
            mode = "metadata_only" if size > 5_000_000 else "sampled_or_referenced"
            notes.append({"filename": name, "path": str(path), "size_bytes": size, "handling": mode})
        else:
            notes.append({"filename": name, "path": str(path), "handling": "not_found"})
    return notes


def make_attachments(root: Path) -> dict[str, Path]:
    for sub in ["xlsx", "pdf", "docx", "csv", "json", "text", "images"]:
        (root / "attachments" / sub).mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    try:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "tasks"
        sheet.append(["id", "client", "amount", "status"])
        for i in range(1, 8):
            sheet.append([i, f"Клиент {i}", 100000 + i * 17000, "active" if i % 2 else "risk"])
        p = root / "attachments" / "xlsx" / "crm_export.xlsx"
        workbook.save(p)
        workbook.close()
        paths["xlsx1"] = p
        for suffix in ["finance_plan", "calendar_slots", "topic_matrix"]:
            q = root / "attachments" / "xlsx" / f"{suffix}.xlsx"
            shutil.copyfile(p, q)
            paths[suffix] = q
    except Exception:
        pass
    try:
        from reportlab.pdfgen import canvas

        for name, text in {
            "client_thread": "Client thread: SLA is close, prepare answer and Jira ticket.",
            "mephi_excerpt": "MEPhI synthetic campus note: admissions, laboratories, timetable.",
            "process_policy": "Process policy excerpt for supplier approval and document checks.",
        }.items():
            p = root / "attachments" / "pdf" / f"{name}.pdf"
            doc = canvas.Canvas(str(p), invariant=1)
            doc.drawString(72, 760, text)
            doc.save()
            paths[name] = p
    except Exception:
        pass
    try:
        from docx import Document

        for name, text in {
            "meeting_notes": "Итоги встречи: проверить сроки, ответственных и риски.",
            "hr_feedback": "Наблюдение руководителя: прогресс, сильные стороны, следующий шаг.",
        }.items():
            p = root / "attachments" / "docx" / f"{name}.docx"
            doc = Document()
            doc.add_paragraph(text)
            doc.save(p)
            paths[name] = p
    except Exception:
        pass
    for name, rows in {
        "sales_rows": [["region", "revenue"], ["north", "120000"], ["south", "98000"]],
        "rooms": [["room", "capacity"], ["A-402", "12"], ["B-110", "24"]],
    }.items():
        p = root / "attachments" / "csv" / f"{name}.csv"
        p.write_text("\n".join(",".join(row) for row in rows) + "\n", encoding="utf-8")
        paths[name] = p
    p = root / "attachments" / "json" / "jira_payload.json"
    write_json(p, {"project": "SYN", "priority": "high", "labels": ["synthetic", "prompt-radar"]})
    paths["jira_payload"] = p
    p = root / "attachments" / "text" / "email_excerpt.md"
    p.write_text("# Email excerpt\nКлиент просит обновить сроки поставки и стоимость.\n", encoding="utf-8")
    paths["email_excerpt"] = p
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d4948445200000001000000010804000000b51c0c020000000b4944415478da63606000000003000168dd8d2d0000000049454e44ae426082"
    )
    for name in ["screen_status", "room_photo"]:
        p = root / "attachments" / "images" / f"{name}.png"
        p.write_bytes(png)
        paths[name] = p
    return paths


def zip_source(source: Path, destination: Path) -> None:
    if destination.exists():
        destination.unlink()
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_STORED) as archive:
        for path in sorted(source.rglob("*")):
            if path.is_file():
                rel = path.relative_to(source).as_posix()
                info = zipfile.ZipInfo(rel, date_time=FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_STORED
                info.external_attr = 0o100644 << 16
                archive.writestr(info, path.read_bytes())


def build() -> dict[str, Any]:
    random.seed(SEED)
    tokenizer = WhitespaceTokenizer()
    if SOURCE_DIR.exists():
        shutil.rmtree(SOURCE_DIR)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    SOURCE_DIR.mkdir(parents=True)
    attachment_files = make_attachments(SOURCE_DIR)

    use_cases = load_yaml(ROOT / "configs" / "known_use_cases.yaml")["use_cases"]
    categories = {c["id"]: c for c in load_yaml(ROOT / "configs" / "categories.yaml")["categories"]}
    users = [
        {"user_id": f"user_{i:03d}", "department": dept, "role": "synthetic_employee", "metadata": {"synthetic": True}}
        for i, dept in enumerate(
            ["HR", "Finance", "Legal", "Sales", "Procurement", "IT", "Project Management", "Operations", "Marketing", "Administration"] * 4,
            start=1,
        )
    ]
    users.append({"user_id": "tatiana_belyakova", "department": "Administration", "role": "student_office_requester", "metadata": {"display_name": "Татьяна Белякова", "synthetic": True}})

    conversations = []
    for i in range(95):
        uid = users[i % 40]["user_id"]
        conversations.append({"conversation_id": f"conv_{i+1:03d}", "user_id": uid, "created_at": "2026-07-01T09:00:00Z", "title": f"Synthetic work chat {i+1:03d}", "metadata": {"synthetic": True}})
    for i in range(2):
        conversations.append({"conversation_id": f"conv_tatiana_{i+1}", "user_id": "tatiana_belyakova", "created_at": f"2026-07-0{i+2}T10:00:00Z", "title": f"Вопросы НИЯУ МИФИ {i+1}", "metadata": {"synthetic": True}})

    known_counts = {case["id"]: 5 for case in use_cases}
    for case in use_cases[-5:]:
        known_counts[case["id"]] = 4
    emerging_groups = {
        "emerging_group_01": "Заявки на внутреннюю библиотеку промптов",
        "emerging_group_02": "Проверка качества данных в витринах",
        "emerging_group_03": "Офисные IoT и датчики помещений",
        "emerging_group_04": "Онбординг внутренних AI-ассистентов",
    }
    mephi_prompts = [
        "Найди правила подачи документов в НИЯУ МИФИ для магистратуры и кратко перечисли дедлайны.",
        "Проверь расписание лабораторных по кафедре и собери аккуратную сводку для группы.",
        "Подскажи, как студенту оформить общежитие МИФИ и какие контакты нужны.",
        "Собери ближайшие мероприятия и олимпиады МИФИ для школьников, без рекламы.",
        "Найди, где в кампусе находятся учебные лаборатории и какие корпуса указать первокурснику.",
    ]
    noise_prompts = [
        "привет", "проверь", "42 17 999", "asdf qwer zzz", "что-нибудь быстро",
        "можно красиво но без смысла", "сегодня чай или отчёт", "test test", "давай потом",
        "почему цифра семь пахнет средой", "hello проверь это самое", "сделай лучше всё",
        "не делай и сделай одновременно", "план", "??! 123 abc", "мне нужен документ наверное",
        "как открыть окно если оно уже открыто", "random unique topic about moon invoices",
        "не знаю спроси у системы", "поставь задачу без задачи", "проверка связи", "ping",
        "слишком общий запрос про бизнес", "переведи но не переводи", "черновик чего-нибудь",
    ]

    specs: list[dict[str, Any]] = []
    for case in use_cases:
        for variant in range(known_counts[case["id"]]):
            specs.append({"group": "known", "use_case": case, "variant": variant})
    for i, prompt in enumerate(mephi_prompts):
        specs.append({"group": "mephi", "prompt": prompt, "variant": i})
    for gid, name in emerging_groups.items():
        for variant in range(5):
            specs.append({"group": "emerging", "cluster_id": gid, "cluster_name": name, "variant": variant})
    for i, prompt in enumerate(noise_prompts):
        specs.append({"group": "noise", "prompt": prompt, "variant": i})
    if len(specs) != 200:
        raise ValueError(len(specs))

    attach_plan = [
        ("xlsx1", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("finance_plan", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("calendar_slots", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("topic_matrix", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("client_thread", "application/pdf"),
        ("mephi_excerpt", "application/pdf"),
        ("process_policy", "application/pdf"),
        ("meeting_notes", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("hr_feedback", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("sales_rows", "text/csv"),
        ("rooms", "text/csv"),
        ("jira_payload", "application/json"),
        ("email_excerpt", "text/markdown"),
        ("screen_status", "image/png"),
        ("room_photo", "image/png"),
    ]
    attachment_indexes = set(range(0, 150, 10))
    runs: list[dict[str, Any]] = []
    messages: list[dict[str, Any]] = []
    events: list[dict[str, Any]] = []
    attachments: list[dict[str, Any]] = []
    ground_truth: list[dict[str, Any]] = []
    token_rows: list[dict[str, Any]] = []
    conv_seq = defaultdict(int)
    statuses = ["completed"] * 188 + ["failed"] * 4 + ["partial"] * 4 + ["cancelled"] * 4
    tool_by_system = {"Email": "email.search", "CRM": "crm.lookup", "Jira": "jira.search", "Calendar": "calendar.find_slot", "Excel": "excel.export", "ИСУП": "isup.ticket", "Confluence": "confluence.search"}

    for idx, spec in enumerate(specs):
        run_id = f"run_{idx+1:03d}"
        if spec["group"] == "mephi":
            conv_id = f"conv_tatiana_{1 + idx % 2}"
        else:
            conv_id = conversations[idx % 95]["conversation_id"]
        started = datetime(2026, 7, 1, 9, tzinfo=timezone.utc) + timedelta(minutes=idx * 17)
        finished = started + timedelta(minutes=3 + idx % 11)
        runs.append({"run_id": run_id, "conversation_id": conv_id, "status": statuses[idx], "started_at": started.isoformat().replace("+00:00", "Z"), "finished_at": finished.isoformat().replace("+00:00", "Z"), "metadata": {"synthetic": True, "generator_group": spec["group"]}})
        bucket, target_tokens = bucket_for_index(idx)
        style = ["plain", "xml", "json", "mixed", "typo"][idx % 5]
        if spec["group"] == "known":
            case = spec["use_case"]
            if target_tokens <= 20:
                base = f"{case['id']} запрос"
            elif target_tokens <= 100:
                base = f"{case['name']}: {case['examples'][0]}. Нужен краткий результат."
            else:
                base = f"{case['name']}: {case['description']} Нужен результат: {case['expected_outcome']}. Вариант {spec['variant'] + 1}, система {', '.join(case.get('systems', []))}."
            use_case_ids = [case["id"]]
            category_ids = list(case["category_ids"])
            discovery = "known"
            cluster_id: str | int | None = None
            cluster_name = None
        elif spec["group"] == "mephi":
            base = spec["prompt"]
            use_case_ids = []
            category_ids = []
            discovery = "emerging"
            cluster_id = "emerging_mephi"
            cluster_name = "Запросы о НИЯУ МИФИ"
        elif spec["group"] == "emerging":
            actions = ["собери", "проверь", "сравни", "подготовь", "разложи по приоритетам"]
            base = f"{actions[spec['variant'] % len(actions)].capitalize()} материалы по теме: {spec['cluster_name']}. Нужен рабочий список действий и рисков, пример {spec['variant'] + 1}."
            use_case_ids = []
            category_ids = []
            discovery = "emerging"
            cluster_id = spec["cluster_id"]
            cluster_name = spec["cluster_name"]
        else:
            base = spec["prompt"]
            use_case_ids = []
            category_ids = []
            discovery = "unresolved"
            cluster_id = -1
            cluster_name = None
        content = make_prompt(base, target_tokens, tokenizer, style)
        actual_tokens = tokenizer.count_tokens(content)
        if actual_tokens != target_tokens:
            raise ValueError((run_id, actual_tokens, target_tokens))
        conv_seq[conv_id] += 1
        message_id = f"msg_{idx+1:03d}_u"
        attachment_ids: list[str] = []
        if idx in attachment_indexes:
            plan_index = sorted(attachment_indexes).index(idx)
            key, media_type = attach_plan[plan_index]
            file_path = attachment_files[key]
            attachment_id = f"att_{plan_index+1:03d}"
            rel = file_path.relative_to(SOURCE_DIR).as_posix()
            attachment_ids.append(attachment_id)
            attachments.append({"attachment_id": attachment_id, "message_id": message_id, "filename": file_path.name, "path": rel, "media_type": media_type, "size_bytes": file_path.stat().st_size, "sha256": sha256_file(file_path), "metadata": {"synthetic": True}})
        messages.append({"message_id": message_id, "conversation_id": conv_id, "run_id": run_id, "role": "user", "content": content, "sequence_number": conv_seq[conv_id], "created_at": started.isoformat().replace("+00:00", "Z"), "attachment_ids": attachment_ids, "metadata": {"style": style}})
        if idx % 6 == 0:
            conv_seq[conv_id] += 1
            messages.append({"message_id": f"msg_{idx+1:03d}_a", "conversation_id": conv_id, "run_id": run_id, "role": "assistant", "content": "Принято, подготовлю результат по указанным ограничениям.", "sequence_number": conv_seq[conv_id], "created_at": (started + timedelta(minutes=1)).isoformat().replace("+00:00", "Z"), "attachment_ids": []})
        tool_names = []
        if spec["group"] == "known":
            tool_names = [tool_by_system[s] for s in spec["use_case"].get("systems", []) if s in tool_by_system]
        if tool_names:
            for tool_index, tool_name in enumerate(tool_names[:2], start=1):
                events.append({"event_id": f"evt_{idx+1:03d}_{tool_index}", "conversation_id": conv_id, "run_id": run_id, "message_id": message_id, "event_type": "tool_call", "status": "success", "sequence_number": tool_index, "created_at": (started + timedelta(seconds=30 * tool_index)).isoformat().replace("+00:00", "Z"), "tool_name": tool_name, "tool_call_id": f"call_{idx+1:03d}_{tool_index}", "payload": {"synthetic": True}})
        ground_truth.append({"run_id": run_id, "synthetic": True, "expected_goal": base, "expected_category_ids": category_ids, "expected_use_case_ids": use_case_ids, "expected_discovery_status": discovery, "expected_cluster_id": cluster_id, "expected_cluster_name": cluster_name, "expected_message_ids": [message_id], "expected_has_attachments": bool(attachment_ids), "expected_token_count": actual_tokens, "length_bucket": bucket, "generator_group": spec["group"]})
        token_rows.append({"run_id": run_id, "token_count": actual_tokens, "length_bucket": bucket, "generator_group": spec["group"]})

    manifest = {"schema_version": "1.0", "dataset_id": "prompt_radar_generated_200", "created_at": "2026-07-25T00:00:00Z", "synthetic": True, "seed": SEED, "generator": "scripts/generate_dataset_200.py", "generator_version": "1.0.0", "source_taxonomy_version": "configs/known_use_cases.yaml:1.0", "tokenizer_model": "WhitespaceTokenizer", "tokenizer_revision": tokenizer.version}
    write_json(SOURCE_DIR / "manifest.json", manifest)
    write_jsonl(SOURCE_DIR / "users.jsonl", users)
    write_jsonl(SOURCE_DIR / "conversations.jsonl", conversations)
    write_jsonl(SOURCE_DIR / "runs.jsonl", runs)
    write_jsonl(SOURCE_DIR / "messages.jsonl", messages)
    write_jsonl(SOURCE_DIR / "events.jsonl", events)
    write_jsonl(SOURCE_DIR / "attachments.jsonl", attachments)
    write_json(SOURCE_DIR / "cost_config.json", {"currency": "RUB", "employee_cost_per_hour": 1500, "version": "generated-200"})
    write_jsonl(GROUND_TRUTH_PATH, ground_truth)
    with TOKEN_COUNTS_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["run_id", "token_count", "length_bucket", "generator_group"])
        writer.writeheader()
        writer.writerows(token_rows)
    README_PATH.write_text("# Generated 200-run dataset\n\nSynthetic Prompt Radar dataset. `ground_truth.jsonl` is intentionally outside `dataset.zip`.\n", encoding="utf-8")

    zip_source(SOURCE_DIR, ZIP_PATH)
    bucket_counts = Counter(row["length_bucket"] for row in token_rows)
    group_counts = Counter(item["generator_group"] for item in ground_truth)
    attachments_by_type = Counter(item["media_type"] for item in attachments)
    report = {
        "total_runs": len(runs),
        "run_count": len(runs),
        "known_runs": group_counts["known"],
        "mephi_runs": group_counts["mephi"],
        "emerging_runs": group_counts["emerging"],
        "noise_runs": group_counts["noise"],
        "known_use_case_count": len(use_cases),
        "emerging_group_count": 5,
        "runs_with_attachments": sum(1 for item in ground_truth if item["expected_has_attachments"]),
        "runs_without_attachments": sum(1 for item in ground_truth if not item["expected_has_attachments"]),
        "total_attachment_files": len(attachments),
        "attachments_by_type": dict(attachments_by_type),
        "exact_100k_token_runs": bucket_counts["exact_100000"],
        "exact_100k_run_ids": [row["run_id"] for row in token_rows if row["length_bucket"] == "exact_100000"],
        "length_distribution": dict(bucket_counts),
        "users": len(users),
        "user_count": len(users),
        "conversations": len(conversations),
        "conversation_count": len(conversations),
        "messages": len(messages),
        "message_count": len(messages),
        "user_message_count": sum(1 for msg in messages if msg["role"] == "user"),
        "events": len(events),
        "attachments": len(attachments),
        "known_distribution": dict(Counter(gt["expected_use_case_ids"][0] for gt in ground_truth if gt["generator_group"] == "known")),
        "emerging_group_names": {"emerging_mephi": "Запросы о НИЯУ МИФИ", **emerging_groups},
        "large_source_file_handling": load_source_notes(),
        "validation_passed": False,
    }
    write_json(REPORT_PATH, report)
    return report


if __name__ == "__main__":
    summary = build()
    print(json.dumps({"dataset_zip": str(ZIP_PATH), "runs": summary["run_count"], "attachments": summary["attachments"]}, ensure_ascii=False, indent=2))
