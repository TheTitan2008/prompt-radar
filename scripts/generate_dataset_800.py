"""Generate the deterministic 800-run Prompt Radar dataset."""

from __future__ import annotations

import csv
import hashlib
import json
import random
import shutil
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from statistics import mean, median
from pathlib import Path
from typing import Any

import yaml

from prompt_radar.preprocessing.tokenizer import WhitespaceTokenizer

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "data" / "generated_800"
SOURCE_DIR = OUT_DIR / "source"
ZIP_PATH = OUT_DIR / "dataset.zip"
GROUND_TRUTH_PATH = OUT_DIR / "ground_truth.jsonl"
REPORT_PATH = OUT_DIR / "generation_report.json"
TOKEN_COUNTS_PATH = OUT_DIR / "token_counts.csv"
README_PATH = OUT_DIR / "README.md"
SEED = 20260725
FIXED_ZIP_TIME = (2026, 7, 25, 0, 0, 0)


def load_yaml(path: Path) -> dict[str, Any]:
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def write_json(path: Path, data: Any) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.write_text("".join(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n" for row in rows), encoding="utf-8")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ensure_token_count(text: str, target: int, tokenizer: WhitespaceTokenizer) -> str:
    current = tokenizer.count_tokens(text)
    if current > target:
        text = " ".join(tokenizer.split_tokens(text)[:target])
        current = tokenizer.count_tokens(text)
    if current == target:
        return text
    filler = " ".join(f"ctx{index:05d}" for index in range(target - current))
    return f"{text}\n\n{filler}"


def bucket_for_index(index: int) -> tuple[str, int]:
    if index < 160:
        return "1_20", 12 + index % 8
    if index < 480:
        return "21_100", 55 + index % 30
    if index < 680:
        return "101_500", 260 + index % 90
    if index < 760:
        return "501_3000", 1200 + index % 300
    if index < 792:
        return "3001_10000", 5200 + index % 600
    return "exact_100000", 100000


def make_prompt(base: str, target_tokens: int, tokenizer: WhitespaceTokenizer, style: str) -> str:
    if target_tokens <= 20:
        return ensure_token_count(base, target_tokens, tokenizer)
    wrappers = {
        "plain": base,
        "xml": f"<user_query>{base}</user_query>",
        "json": json.dumps(
            {
                "messages": [
                    {"role": "system", "content": "Внутренний контекст не является целью пользователя."},
                    {"role": "user", "content": f"<context>Рабочая задача</context><user_query>{base}</user_query>"},
                ]
            },
            ensure_ascii=False,
        ),
        "mixed": f"Need help: {base}. Please keep result structured.",
        "typo": base.replace("отчёт", "атчёт").replace("задачу", "задочу").replace("сводку", "сфодку"),
    }
    candidate = wrappers[style]
    if tokenizer.count_tokens(candidate) > target_tokens:
        candidate = base
    return ensure_token_count(candidate, target_tokens, tokenizer)


def deterministic_fraction(*parts: object) -> float:
    raw = ":".join(str(part) for part in (SEED, *parts))
    value = int(hashlib.sha256(raw.encode("utf-8")).hexdigest()[:12], 16)
    return value / float(0xFFFFFFFFFFFF)


def calculate_ai_processing_minutes(
    *,
    run_id: str,
    token_count: int,
    length_bucket: str,
    attachment_media_types: list[str],
    message_count: int,
    tool_count: int,
    generator_group: str,
    status: str,
) -> float:
    jitter = deterministic_fraction(run_id, token_count, generator_group)
    if length_bucket == "exact_100000":
        base = 45 + 110 * jitter
    elif length_bucket == "3001_10000":
        base = 15 + min((token_count - 3000) / 7000, 1) * 30 + 12 * jitter
    elif length_bucket == "501_3000":
        base = 6 + min((token_count - 500) / 2500, 1) * 12 + 5 * jitter
    elif length_bucket == "101_500":
        base = 3 + min((token_count - 100) / 400, 1) * 5 + 3 * jitter
    elif length_bucket == "21_100":
        base = 1 + min((token_count - 20) / 80, 1) * 3 + 1.5 * jitter
    else:
        base = 0.2 + min(token_count / 20, 1) * 1.6 + 0.7 * jitter
        if generator_group == "known":
            base += 0.8

    group_add = {
        "known": 1.0,
        "mephi": 1.4,
        "emerging": 1.8,
        "minor_emerging": 1.0,
        "noise": 0.0,
    }[generator_group]
    base += group_add
    for media_type in attachment_media_types:
        if media_type.endswith("spreadsheetml.sheet"):
            base += 5.5
        elif media_type == "application/pdf":
            base += 5.0
        elif media_type.endswith("wordprocessingml.document"):
            base += 4.0
        elif media_type in {"text/csv", "application/json"}:
            base += 2.0
        elif media_type.startswith("image/"):
            base += 0.8
        else:
            base += 1.0
    base += tool_count * 0.9
    base += max(message_count - 1, 0) * 0.5

    status_factor = {
        "completed": 1.0,
        "partial": 0.72,
        "failed": 0.58,
        "cancelled": 0.42,
    }.get(status, 1.0)
    minutes = base * status_factor
    if generator_group == "noise" and length_bucket in {"1_20", "21_100"}:
        minutes = min(minutes, 2.0)
    if length_bucket == "exact_100000":
        minutes = min(max(minutes, 45), 180)
    return round(max(minutes, 0.2), 2)


def time_bucket(minutes: float) -> str:
    if minutes < 2:
        return "under_2"
    if minutes < 5:
        return "2_5"
    if minutes < 12:
        return "5_12"
    if minutes < 25:
        return "12_25"
    if minutes < 60:
        return "25_60"
    return "60_plus"


def iso_z(value: datetime) -> str:
    return value.isoformat().replace("+00:00", "Z")


def source_notes() -> list[dict[str, Any]]:
    downloads = Path.home() / "Downloads"
    requested = [
        "Темы для генерации датасета.xlsx",
        "17 (2).txt",
        "27_B__abzyn.txt",
        "вариант 1.1 new 2024.docx",
        "Демодень_ДЗ (1).ipynb",
        "ДЗ_демодень_Кирилин_Александр_ipynb_.ipynb",
        "Menschen_A1_2_AB.pdf",
        "Menschen_A1_2_KB.pdf",
        "DEMO_24 (1).txt",
        "DEMO_24 (2).txt",
    ]
    attachment_dir = Path(r"C:\Users\aleks\.codex\attachments\732e0bb4-3010-45a6-ab86-724561fde496")
    requested.extend(["image-1.jpg", "image-2.png"])
    notes: list[dict[str, Any]] = []
    for name in requested:
        path = attachment_dir / name if name.startswith("image-") else downloads / name
        if path.exists():
            size = path.stat().st_size
            mode = "metadata_only" if size > 5_000_000 else "sampled_or_referenced"
            notes.append({"filename": name, "path": str(path), "size_bytes": size, "handling": mode})
        else:
            notes.append({"filename": name, "path": str(path), "handling": "not_found"})
    return notes


def make_attachments(root: Path) -> dict[str, tuple[Path, str]]:
    for sub in ["xlsx", "pdf", "docx", "csv", "json", "text", "images"]:
        (root / "attachments" / sub).mkdir(parents=True, exist_ok=True)
    paths: dict[str, tuple[Path, str]] = {}
    try:
        from openpyxl import Workbook

        for book in range(16):
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "tasks"
            sheet.append(["id", "client", "amount", "status"])
            for i in range(1, 8):
                sheet.append([i, f"Клиент {book + 1}-{i}", 100000 + i * 17000 + book * 1000, "active" if i % 2 else "risk"])
            p = root / "attachments" / "xlsx" / f"workbook_{book + 1:02d}.xlsx"
            workbook.save(p)
            workbook.close()
            paths[f"xlsx_{book}"] = (p, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception:
        pass
    try:
        from reportlab.pdfgen import canvas

        for i in range(12):
            p = root / "attachments" / "pdf" / f"excerpt_{i + 1:02d}.pdf"
            doc = canvas.Canvas(str(p), invariant=1)
            doc.drawString(72, 760, f"Synthetic source excerpt {i + 1}: SLA, MEPhI, supplier process, schedule.")
            doc.save()
            paths[f"pdf_{i}"] = (p, "application/pdf")
    except Exception:
        pass
    try:
        from docx import Document

        for i in range(8):
            p = root / "attachments" / "docx" / f"note_{i + 1:02d}.docx"
            doc = Document()
            doc.add_paragraph(f"Синтетическая заметка {i + 1}: итоги встречи, риски, следующие шаги.")
            doc.save(p)
            paths[f"docx_{i}"] = (p, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except Exception:
        pass
    for i in range(8):
        p = root / "attachments" / "csv" / f"rows_{i + 1:02d}.csv"
        p.write_text("region,revenue\nnorth,120000\nsouth,98000\n", encoding="utf-8")
        paths[f"csv_{i}"] = (p, "text/csv")
    for i in range(4):
        p = root / "attachments" / "json" / f"payload_{i + 1:02d}.json"
        write_json(p, {"project": "SYN", "priority": "high", "index": i + 1})
        paths[f"json_{i}"] = (p, "application/json")
    for i in range(4):
        p = root / "attachments" / "text" / f"excerpt_{i + 1:02d}.md"
        p.write_text(f"# Excerpt {i + 1}\nКлиент просит обновить сроки поставки и стоимость.\n", encoding="utf-8")
        paths[f"text_{i}"] = (p, "text/markdown")
    embedded_png = bytes.fromhex(
        "89504e470d0a1a0a0000000d4948445200000001000000010804000000b51c0c020000000b4944415478da63606000000003000168dd8d2d0000000049454e44ae426082"
    )
    external = [
        Path(r"C:\Users\aleks\.codex\attachments\732e0bb4-3010-45a6-ab86-724561fde496\image-1.jpg"),
        Path(r"C:\Users\aleks\.codex\attachments\732e0bb4-3010-45a6-ab86-724561fde496\image-2.png"),
    ]
    for i in range(8):
        if i < 2 and external[i].exists():
            suffix = external[i].suffix.lower()
            p = root / "attachments" / "images" / f"attached_image_{i + 1:02d}{suffix}"
            shutil.copyfile(external[i], p)
            media = "image/jpeg" if suffix in {".jpg", ".jpeg"} else "image/png"
        else:
            p = root / "attachments" / "images" / f"synthetic_image_{i + 1:02d}.png"
            p.write_bytes(embedded_png)
            media = "image/png"
        paths[f"image_{i}"] = (p, media)
    return paths


def zip_source(source: Path, destination: Path) -> None:
    destination.unlink(missing_ok=True)
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_STORED) as archive:
        for path in sorted(source.rglob("*")):
            if path.is_file():
                rel = path.relative_to(source).as_posix()
                info = zipfile.ZipInfo(rel, date_time=FIXED_ZIP_TIME)
                info.compress_type = zipfile.ZIP_STORED
                info.external_attr = 0o100644 << 16
                archive.writestr(info, path.read_bytes())


def named_users() -> list[dict[str, Any]]:
    roster = [
        ("aleksei_smirnov", "Алексей Смирнов", "Sales", 106),
        ("elena_petrova", "Елена Петрова", "Operations", 94),
        ("dmitry_ivanov", "Дмитрий Иванов", "IT", 82),
        ("marina_sokolova", "Марина Соколова", "Finance", 70),
        ("sergey_kuznetsov", "Сергей Кузнецов", "Project Management", 58),
        ("olga_morozova", "Ольга Морозова", "HR", 52),
        ("pavel_volkov", "Павел Волков", "Procurement", 48),
        ("natalia_lebedeva", "Наталья Лебедева", "Legal", 44),
        ("igor_fedorov", "Игорь Федоров", "Marketing", 40),
        ("anna_novikova", "Анна Новикова", "Administration", 36),
        ("roman_egorov", "Роман Егоров", "Sales", 32),
        ("ekaterina_orlova", "Екатерина Орлова", "Finance", 28),
        ("mikhail_popov", "Михаил Попов", "IT", 24),
        ("irina_vasilieva", "Ирина Васильева", "Operations", 20),
        ("tatiana_belyakova", "Татьяна Белякова", "Administration", 20),
        ("viktor_zaitsev", "Виктор Зайцев", "Project Management", 16),
        ("svetlana_krylova", "Светлана Крылова", "HR", 12),
        ("boris_maksimov", "Борис Максимов", "Procurement", 8),
        ("yulia_belova", "Юлия Белова", "Marketing", 6),
        ("nikita_andreev", "Никита Андреев", "Legal", 4),
    ]
    if sum(item[3] for item in roster) != 800:
        raise ValueError("User distribution must sum to 800")
    return [
        {
            "user_id": user_id,
            "display_name": display_name,
            "department": department,
            "role": "synthetic_employee",
            "metadata": {"synthetic": True, "planned_run_count": planned},
        }
        for user_id, display_name, department, planned in roster
    ]


def user_distribution(users: list[dict[str, Any]]) -> list[str]:
    distribution: list[str] = []
    for user in users:
        if user["user_id"] == "tatiana_belyakova":
            continue
        distribution.extend([user["user_id"]] * int(user["metadata"]["planned_run_count"]))
    if len(distribution) != 780:
        raise ValueError("Non-MEPhI user distribution must contain 780 runs")
    return distribution


def conversations_for_users(users: list[dict[str, Any]]) -> list[dict[str, Any]]:
    conversations: list[dict[str, Any]] = []
    for user in users:
        planned = int(user["metadata"]["planned_run_count"])
        count = max(1, min(18, (planned + 5) // 6))
        if user["user_id"] == "tatiana_belyakova":
            count = 4
        for index in range(count):
            conversation_id = f"conv_{user['user_id']}_{index + 1:02d}"
            conversations.append(
                {
                    "conversation_id": conversation_id,
                    "user_id": user["user_id"],
                    "owner_user_id": user["user_id"],
                    "created_at": f"2026-07-{1 + index % 20:02d}T09:00:00Z",
                    "title": f"Рабочий чат: {user['display_name']} #{index + 1}",
                    "metadata": {"synthetic": True},
                }
            )
    return conversations


def build_specs(use_cases: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, str], dict[str, str]]:
    known_counts = {case["id"]: 20 for case in use_cases}
    for case in use_cases[-5:]:
        known_counts[case["id"]] = 16
    major_groups = {
        "emerging_group_01": "Заявки на внутреннюю библиотеку промптов",
        "emerging_group_02": "Проверка качества данных в витринах",
        "emerging_group_03": "Офисные IoT и датчики помещений",
        "emerging_group_04": "Онбординг внутренних AI-ассистентов",
    }
    minor_groups = {
        "minor_group_01": "Микроопросы удовлетворенности сотрудников",
        "minor_group_02": "Сверка витрин учебных материалов",
        "minor_group_03": "Контроль локальных справочников оборудования",
        "minor_group_04": "Черновики внутренних объявлений",
        "minor_group_05": "Проверка маленьких списков закупок",
    }
    specs: list[dict[str, Any]] = []
    for case in use_cases:
        for variant in range(known_counts[case["id"]]):
            specs.append({"group": "known", "use_case": case, "variant": variant})
    mephi_bases = [
        "Найди правила подачи документов в НИЯУ МИФИ для магистратуры и кратко перечисли дедлайны.",
        "Проверь расписание лабораторных по кафедре и собери аккуратную сводку для группы.",
        "Подскажи, как студенту оформить общежитие МИФИ и какие контакты нужны.",
        "Собери ближайшие мероприятия и олимпиады МИФИ для школьников, без рекламы.",
        "Найди, где в кампусе находятся учебные лаборатории и какие корпуса указать первокурснику.",
    ]
    for i in range(20):
        specs.append({"group": "mephi", "prompt": f"{mephi_bases[i % len(mephi_bases)]} Версия запроса {i + 1}.", "variant": i})
    for gid, name in major_groups.items():
        for variant in range(20):
            specs.append({"group": "emerging", "cluster_id": gid, "cluster_name": name, "variant": variant})
    for gid, name in minor_groups.items():
        for variant in range(12):
            specs.append({"group": "minor_emerging", "cluster_id": gid, "cluster_name": name, "variant": variant})
    noise = [
        "привет", "проверь", "42 17 999", "asdf qwer zzz", "что-нибудь быстро",
        "можно красиво но без смысла", "сегодня чай или отчёт", "test test", "давай потом",
        "почему цифра семь пахнет средой", "hello проверь это самое", "сделай лучше всё",
        "не делай и сделай одновременно", "план", "??! 123 abc", "мне нужен документ наверное",
        "как открыть окно если оно уже открыто", "random unique topic about moon invoices",
        "не знаю спроси у системы", "поставь задачу без задачи", "проверка связи", "ping",
        "слишком общий запрос про бизнес", "переведи но не переводи", "черновик чего-нибудь",
        "1234567890", "ок", "да", "нет", "потом", "синий бюджет?", "обнови штуку",
        "без контекста сделай", "final final draft", "забыл что хотел", "сравни неизвестно что",
        "просто список", "проверка голоса", "календарь CRM Jira всё сразу без цели", "..."
    ]
    for i in range(40):
        specs.append({"group": "noise", "prompt": noise[i], "variant": i})
    if len(specs) != 800:
        raise ValueError(f"Expected 800 specs, got {len(specs)}")
    return specs, major_groups, minor_groups


def build() -> dict[str, Any]:
    random.seed(SEED)
    tokenizer = WhitespaceTokenizer()
    if SOURCE_DIR.exists():
        shutil.rmtree(SOURCE_DIR)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    SOURCE_DIR.mkdir(parents=True)
    attachment_files = make_attachments(SOURCE_DIR)
    attach_keys = sorted(attachment_files)

    use_cases = load_yaml(ROOT / "configs" / "known_use_cases.yaml")["use_cases"]
    specs, major_groups, minor_groups = build_specs(use_cases)
    users = named_users()
    user_by_id = {user["user_id"]: user for user in users}
    non_mephi_user_queue = user_distribution(users)
    non_mephi_cursor = 0
    conversations = conversations_for_users(users)
    conversations_by_user: dict[str, list[str]] = defaultdict(list)
    for conversation in conversations:
        conversations_by_user[conversation["user_id"]].append(conversation["conversation_id"])
    run_counter_by_user: Counter[str] = Counter()

    attachment_indexes = set(range(0, 600, 10))
    statuses = ["completed"] * 752 + ["failed"] * 16 + ["partial"] * 16 + ["cancelled"] * 16
    tool_by_system = {"Email": "email.search", "CRM": "crm.lookup", "Jira": "jira.search", "Calendar": "calendar.find_slot", "Excel": "excel.export", "ИСУП": "isup.ticket", "Confluence": "confluence.search"}
    runs: list[dict[str, Any]] = []
    messages: list[dict[str, Any]] = []
    events: list[dict[str, Any]] = []
    attachments: list[dict[str, Any]] = []
    ground_truth: list[dict[str, Any]] = []
    token_rows: list[dict[str, Any]] = []
    conv_seq = defaultdict(int)
    styles = ["plain", "xml", "json", "mixed", "typo"]
    actions = ["собери", "проверь", "сравни", "подготовь", "разложи по приоритетам", "сделай короткую сводку"]

    for idx, spec in enumerate(specs):
        run_id = f"run_{idx + 1:03d}"
        if spec["group"] == "mephi":
            user_id = "tatiana_belyakova"
        else:
            user_id = non_mephi_user_queue[non_mephi_cursor]
            non_mephi_cursor += 1
        user = user_by_id[user_id]
        user_conversations = conversations_by_user[user_id]
        conv_id = user_conversations[run_counter_by_user[user_id] % len(user_conversations)]
        run_counter_by_user[user_id] += 1
        started = datetime(2026, 7, 1, 9, tzinfo=timezone.utc) + timedelta(minutes=idx * 11)
        status = statuses[idx]
        bucket, target_tokens = bucket_for_index(idx)
        style = styles[idx % len(styles)]
        if spec["group"] == "known":
            case = spec["use_case"]
            if target_tokens <= 20:
                base = f"{case['id']} запрос {spec['variant'] + 1}"
            elif target_tokens <= 100:
                base = f"{case['name']}: {case['examples'][0]}. Нужен краткий результат, вариант {spec['variant'] + 1}."
            else:
                base = f"{case['name']}: {case['description']} Нужен результат: {case['expected_outcome']}. Вариант {spec['variant'] + 1}, система {', '.join(case.get('systems', []))}."
            use_case_ids = [case["id"]]
            category_ids = list(case["category_ids"])
            discovery = "known"
            cluster_id: str | int | None = None
            cluster_name = None
        elif spec["group"] == "mephi":
            base = spec["prompt"]
            use_case_ids = []
            category_ids = []
            discovery = "emerging"
            cluster_id = "emerging_mephi"
            cluster_name = "Запросы о НИЯУ МИФИ"
        elif spec["group"] in {"emerging", "minor_emerging"}:
            base = f"{actions[spec['variant'] % len(actions)].capitalize()} материалы по теме: {spec['cluster_name']}. Нужен рабочий список действий и рисков, пример {spec['variant'] + 1}."
            use_case_ids = []
            category_ids = []
            discovery = "emerging"
            cluster_id = spec["cluster_id"]
            cluster_name = spec["cluster_name"]
        else:
            base = spec["prompt"]
            use_case_ids = []
            category_ids = []
            discovery = "unresolved"
            cluster_id = -1
            cluster_name = None
        content = make_prompt(base, target_tokens, tokenizer, style)
        actual_tokens = tokenizer.count_tokens(content)
        if actual_tokens != target_tokens:
            raise ValueError((run_id, actual_tokens, target_tokens))
        conv_seq[conv_id] += 1
        message_id = f"msg_{idx + 1:03d}_u"
        attachment_ids: list[str] = []
        if idx in attachment_indexes:
            plan_index = sorted(attachment_indexes).index(idx)
            key = attach_keys[plan_index % len(attach_keys)]
            file_path, media_type = attachment_files[key]
            attachment_id = f"att_{plan_index + 1:03d}"
            rel = file_path.relative_to(SOURCE_DIR).as_posix()
            attachment_ids.append(attachment_id)
            attachments.append({"attachment_id": attachment_id, "message_id": message_id, "filename": file_path.name, "path": rel, "media_type": media_type, "size_bytes": file_path.stat().st_size, "sha256": sha256_file(file_path), "metadata": {"synthetic": True}})
        tool_names = []
        if spec["group"] == "known":
            tool_names = [tool_by_system[s] for s in spec["use_case"].get("systems", []) if s in tool_by_system]
        assistant_message_count = 1 if idx % 6 == 0 else 0
        attachment_media_types = [
            item["media_type"]
            for item in attachments
            if item["attachment_id"] in attachment_ids
        ]
        ai_processing_minutes = calculate_ai_processing_minutes(
            run_id=run_id,
            token_count=actual_tokens,
            length_bucket=bucket,
            attachment_media_types=attachment_media_types,
            message_count=1 + assistant_message_count,
            tool_count=min(len(tool_names), 2),
            generator_group=spec["group"],
            status=status,
        )
        finished = started + timedelta(minutes=ai_processing_minutes)
        runs.append(
            {
                "run_id": run_id,
                "conversation_id": conv_id,
                "user_id": user_id,
                "status": status,
                "started_at": iso_z(started),
                "finished_at": iso_z(finished),
                "metadata": {
                    "synthetic": True,
                    "generator_group": spec["group"],
                    "ai_processing_minutes": ai_processing_minutes,
                    "ai_processing_minutes_source": "synthetic_complexity_estimate",
                    "ai_processing_minutes_evidence_level": "E0",
                },
            }
        )
        messages.append({"message_id": message_id, "conversation_id": conv_id, "run_id": run_id, "sender_user_id": user_id, "role": "user", "content": content, "sequence_number": conv_seq[conv_id], "created_at": iso_z(started), "attachment_ids": attachment_ids, "metadata": {"style": style}})
        if idx % 6 == 0:
            conv_seq[conv_id] += 1
            messages.append({"message_id": f"msg_{idx + 1:03d}_a", "conversation_id": conv_id, "run_id": run_id, "role": "assistant", "content": "Принято, подготовлю результат по указанным ограничениям.", "sequence_number": conv_seq[conv_id], "created_at": iso_z(started + timedelta(seconds=min(ai_processing_minutes * 60 * 0.5, max(ai_processing_minutes * 60 - 1, 0)))), "attachment_ids": []})
        for tool_index, tool_name in enumerate(tool_names[:2], start=1):
            event_offset = max(1.0, min(ai_processing_minutes * 60 * (0.2 + 0.2 * tool_index), ai_processing_minutes * 60 - 1))
            events.append({"event_id": f"evt_{idx + 1:03d}_{tool_index}", "conversation_id": conv_id, "run_id": run_id, "message_id": message_id, "event_type": "tool_call", "status": "success", "sequence_number": tool_index, "created_at": iso_z(started + timedelta(seconds=event_offset)), "tool_name": tool_name, "tool_call_id": f"call_{idx + 1:03d}_{tool_index}", "payload": {"synthetic": True}})
        ground_truth.append({"run_id": run_id, "synthetic": True, "expected_goal": base, "expected_category_ids": category_ids, "expected_use_case_ids": use_case_ids, "expected_discovery_status": discovery, "expected_cluster_id": cluster_id, "expected_cluster_name": cluster_name, "expected_message_ids": [message_id], "expected_has_attachments": bool(attachment_ids), "expected_token_count": actual_tokens, "expected_user_id": user_id, "expected_user_name": user["display_name"], "expected_ai_processing_minutes": ai_processing_minutes, "expected_ai_processing_minutes_source": "synthetic_complexity_estimate", "length_bucket": bucket, "generator_group": spec["group"]})
        token_rows.append({"run_id": run_id, "token_count": actual_tokens, "length_bucket": bucket, "generator_group": spec["group"]})

    if non_mephi_cursor != len(non_mephi_user_queue):
        raise ValueError("Non-MEPhI user distribution was not fully consumed")
    if run_counter_by_user["tatiana_belyakova"] != 20:
        raise ValueError("Tatiana Belyakova must own exactly 20 MEPhI runs")

    write_json(SOURCE_DIR / "manifest.json", {"schema_version": "1.0", "dataset_id": "prompt_radar_generated_800", "created_at": "2026-07-25T00:00:00Z", "synthetic": True, "seed": SEED, "generator": "scripts/generate_dataset_800.py", "generator_version": "1.0.0", "source_taxonomy_version": "configs/known_use_cases.yaml:1.0", "tokenizer_model": "WhitespaceTokenizer", "tokenizer_revision": tokenizer.version})
    write_jsonl(SOURCE_DIR / "users.jsonl", users)
    write_jsonl(SOURCE_DIR / "conversations.jsonl", conversations)
    write_jsonl(SOURCE_DIR / "runs.jsonl", runs)
    write_jsonl(SOURCE_DIR / "messages.jsonl", messages)
    write_jsonl(SOURCE_DIR / "events.jsonl", events)
    write_jsonl(SOURCE_DIR / "attachments.jsonl", attachments)
    write_json(SOURCE_DIR / "cost_config.json", {"currency": "RUB", "employee_cost_per_hour": 1500, "version": "generated-800"})
    write_jsonl(GROUND_TRUTH_PATH, ground_truth)
    with TOKEN_COUNTS_PATH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["run_id", "token_count", "length_bucket", "generator_group"])
        writer.writeheader()
        writer.writerows(token_rows)
    README_PATH.write_text(
        "# Generated 800-run dataset\n\n"
        "Synthetic Prompt Radar dataset. `ground_truth.jsonl` is intentionally "
        "outside `dataset.zip`.\n\n"
        "`run.metadata.ai_processing_minutes` is a deterministic E0 synthetic "
        "complexity estimate of AI processing time, not measured human prompt "
        "authoring time.\n",
        encoding="utf-8",
    )
    zip_source(SOURCE_DIR, ZIP_PATH)

    bucket_counts = Counter(row["length_bucket"] for row in token_rows)
    group_counts = Counter(item["generator_group"] for item in ground_truth)
    runs_by_user_id = Counter(run["user_id"] for run in runs)
    runs_by_user = {
        user["display_name"]: runs_by_user_id[user["user_id"]]
        for user in users
    }
    ai_values = [
        float(run["metadata"]["ai_processing_minutes"])
        for run in runs
    ]
    interval_mismatches = 0
    for run in runs:
        started_at = datetime.fromisoformat(run["started_at"].replace("Z", "+00:00"))
        finished_at = datetime.fromisoformat(run["finished_at"].replace("Z", "+00:00"))
        actual = (finished_at - started_at).total_seconds() / 60
        expected = float(run["metadata"]["ai_processing_minutes"])
        if abs(actual - expected) > (1 / 60):
            interval_mismatches += 1
    report = {
        "total_runs": len(runs),
        "run_count": len(runs),
        "known_runs": group_counts["known"],
        "mephi_runs": group_counts["mephi"],
        "emerging_runs": group_counts["emerging"],
        "minor_emerging_runs": group_counts["minor_emerging"],
        "noise_runs": group_counts["noise"],
        "noise_block_runs": group_counts["minor_emerging"] + group_counts["noise"],
        "known_use_case_count": len(use_cases),
        "emerging_group_count": 10,
        "major_emerging_group_count": 4,
        "minor_emerging_group_count": 5,
        "runs_with_attachments": sum(1 for item in ground_truth if item["expected_has_attachments"]),
        "runs_without_attachments": sum(1 for item in ground_truth if not item["expected_has_attachments"]),
        "total_attachment_files": len(attachments),
        "attachments_by_type": dict(Counter(item["media_type"] for item in attachments)),
        "exact_100k_token_runs": bucket_counts["exact_100000"],
        "exact_100k_run_ids": [row["run_id"] for row in token_rows if row["length_bucket"] == "exact_100000"],
        "length_distribution": dict(bucket_counts),
        "unique_named_users": len(users),
        "runs_by_user": runs_by_user,
        "min_runs_per_user": min(runs_by_user.values()),
        "max_runs_per_user": max(runs_by_user.values()),
        "mean_runs_per_user": round(mean(runs_by_user.values()), 2),
        "median_runs_per_user": median(runs_by_user.values()),
        "ai_processing_minutes_min": min(ai_values),
        "ai_processing_minutes_max": max(ai_values),
        "ai_processing_minutes_mean": round(mean(ai_values), 2),
        "ai_processing_minutes_median": median(ai_values),
        "ai_processing_minutes_by_complexity": dict(Counter(time_bucket(value) for value in ai_values)),
        "ai_processing_interval_mismatch_count": interval_mismatches,
        "mephi_runs_owned_by_tatiana_belyakova": sum(
            1
            for item in ground_truth
            if item["generator_group"] == "mephi"
            and item["expected_user_id"] == "tatiana_belyakova"
        ),
        "users": len(users),
        "user_count": len(users),
        "conversations": len(conversations),
        "conversation_count": len(conversations),
        "messages": len(messages),
        "message_count": len(messages),
        "user_message_count": sum(1 for msg in messages if msg["role"] == "user"),
        "events": len(events),
        "attachments": len(attachments),
        "known_distribution": dict(Counter(gt["expected_use_case_ids"][0] for gt in ground_truth if gt["generator_group"] == "known")),
        "emerging_group_names": {"emerging_mephi": "Запросы о НИЯУ МИФИ", **major_groups, **minor_groups},
        "large_source_file_handling": source_notes(),
        "validation_passed": False,
    }
    write_json(REPORT_PATH, report)
    return report


if __name__ == "__main__":
    summary = build()
    print(json.dumps({"dataset_zip": str(ZIP_PATH), "runs": summary["run_count"], "attachments": summary["attachments"]}, ensure_ascii=False, indent=2))
